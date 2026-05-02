"""
FitGuard AI – Yazılım Teknik Dokümantasyonu (DOCX üreteci)

Koddaki gerçek davranışı (backend/app/*, frontend/lib/*) doğrudan baz alır.
Çıktı: .../sunum/FitGuardAI_Yazilim_Dokumantasyonu.docx
"""

from __future__ import annotations

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor


# --- Renk paleti (sunumlarla uyumlu) -----------------------------------------
NAVY = RGBColor(0x0B, 0x17, 0x29)
MINT = RGBColor(0x10, 0xB9, 0x81)
SLATE900 = RGBColor(0x0F, 0x17, 0x2A)
SLATE700 = RGBColor(0x33, 0x41, 0x55)
SLATE500 = RGBColor(0x64, 0x74, 0x8B)
BG_LIGHT = "F1F5F9"
BG_CODE = "0B1729"


def shade_cell(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def set_cell_border(cell, color: str = "E2E8F0", sz: int = 6) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(sz))
        b.set(qn("w:color"), color)
        tc_borders.append(b)
    tc_pr.append(tc_borders)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "Calibri"
        if level == 0:
            run.font.size = Pt(28)
            run.font.color.rgb = NAVY
        elif level == 1:
            run.font.size = Pt(18)
            run.font.color.rgb = NAVY
        elif level == 2:
            run.font.size = Pt(14)
            run.font.color.rgb = SLATE900
        else:
            run.font.size = Pt(12)
            run.font.color.rgb = SLATE700
        run.bold = True


def add_paragraph(doc: Document, text: str, size: int = 11, color: RGBColor = SLATE700, bold: bool = False, italic: bool = False) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.bold = bold
    r.italic = italic


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(item)
        r.font.name = "Calibri"
        r.font.size = Pt(11)
        r.font.color.rgb = SLATE700


def add_code_block(doc: Document, code: str, caption: str | None = None) -> None:
    if caption:
        cp = doc.add_paragraph()
        cr = cp.add_run(caption)
        cr.font.name = "Calibri"
        cr.font.size = Pt(10)
        cr.italic = True
        cr.font.color.rgb = SLATE500

    tbl = doc.add_table(rows=1, cols=1)
    tbl.autofit = False
    cell = tbl.rows[0].cells[0]
    cell.width = Cm(16)
    shade_cell(cell, BG_CODE)
    set_cell_border(cell, color="0B1729", sz=4)
    cell.paragraphs[0].clear()
    for i, line in enumerate(code.splitlines() or [""]):
        if i == 0:
            p = cell.paragraphs[0]
        else:
            p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        r = p.add_run(line if line else "\u00a0")
        r.font.name = "Consolas"
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
    doc.add_paragraph()


def add_callout(doc: Document, title: str, body: str) -> None:
    tbl = doc.add_table(rows=1, cols=1)
    tbl.autofit = False
    cell = tbl.rows[0].cells[0]
    cell.width = Cm(16)
    shade_cell(cell, BG_LIGHT)
    set_cell_border(cell, color="10B981", sz=12)
    cell.paragraphs[0].clear()
    p1 = cell.paragraphs[0]
    r1 = p1.add_run(title)
    r1.bold = True
    r1.font.name = "Calibri"
    r1.font.size = Pt(11)
    r1.font.color.rgb = MINT
    p2 = cell.add_paragraph()
    r2 = p2.add_run(body)
    r2.font.name = "Calibri"
    r2.font.size = Pt(10)
    r2.font.color.rgb = SLATE700
    doc.add_paragraph()


def add_kv_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    tbl = doc.add_table(rows=len(rows), cols=2)
    tbl.autofit = False
    for i, (k, v) in enumerate(rows):
        c1, c2 = tbl.rows[i].cells
        c1.width = Cm(5)
        c2.width = Cm(11)
        shade_cell(c1, "0B1729")
        shade_cell(c2, "F8FAFC")
        set_cell_border(c1, color="0B1729", sz=4)
        set_cell_border(c2, color="E2E8F0", sz=4)
        c1.paragraphs[0].clear()
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(k)
        r1.bold = True
        r1.font.name = "Calibri"
        r1.font.size = Pt(10)
        r1.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        c2.paragraphs[0].clear()
        p2 = c2.paragraphs[0]
        r2 = p2.add_run(v)
        r2.font.name = "Calibri"
        r2.font.size = Pt(10)
        r2.font.color.rgb = SLATE700
    doc.add_paragraph()


# === Doküman başlangıcı ======================================================

doc = Document()

for section in doc.sections:
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

# --- Kapak sayfası -----------------------------------------------------------
cover = doc.add_paragraph()
cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = cover.add_run("FITGUARD AI\n")
r.font.name = "Calibri"
r.font.size = Pt(32)
r.bold = True
r.font.color.rgb = NAVY

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Yazılım Teknik Dokümantasyonu")
r.font.name = "Calibri"
r.font.size = Pt(18)
r.italic = True
r.font.color.rgb = MINT

tag = doc.add_paragraph()
tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = tag.add_run(
    "Yapay Zeka Kişiselleştirilmiş Gerçek Zamanlı Egzersiz Form Analiz Sistemi"
)
r.font.name = "Calibri"
r.font.size = Pt(12)
r.font.color.rgb = SLATE500

doc.add_paragraph()
doc.add_paragraph()
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = meta.add_run("TEKNOFEST 2026 — Sağlık, Sağlıklı Yaşam ve Spor Kategorisi")
r.font.name = "Calibri"
r.font.size = Pt(11)
r.font.color.rgb = SLATE700

doc.add_page_break()

# --- İçindekiler -------------------------------------------------------------
add_heading(doc, "İçindekiler", level=1)
toc = [
    "1. Genel Bakış",
    "2. Sistem Mimarisi",
    "3. Teknoloji Yığını",
    "4. Proje Dizin Yapısı",
    "5. Backend Bileşenleri",
    "   5.1 Giriş Noktası (main.py)",
    "   5.2 Konfigürasyon Katmanı",
    "   5.3 Egzersiz Tanımları",
    "   5.4 Çekirdek Modüller (core/)",
    "   5.5 Yapay Zeka Katmanı (ai/)",
    "   5.6 Servisler (services/)",
    "   5.7 WebSocket İşleyicisi",
    "   5.8 Veri Modelleri ve Kalıcılık",
    "6. Frontend Bileşenleri",
    "7. Uçtan Uca Kullanıcı Akışı",
    "8. Güvenlik, Gizlilik ve KVKK",
    "9. Geliştirme ve Çalıştırma",
]
for line in toc:
    p = doc.add_paragraph()
    r = p.add_run(line)
    r.font.name = "Calibri"
    r.font.size = Pt(11)
    r.font.color.rgb = SLATE700

doc.add_page_break()

# === 1. Genel Bakış ==========================================================
add_heading(doc, "1. Genel Bakış", level=1)
add_paragraph(
    doc,
    "FitGuard AI, standart bir web kamerası üzerinden kullanıcının egzersiz "
    "formunu gerçek zamanlı analiz eden, tıbbi bir dil modeli (MedGemma) ile "
    "kişisel kalibrasyonu birleştiren ve tüm veriyi cihaz üzerinde işleyen bir "
    "sistemdir. Çözüm üç katmandan oluşur: (i) tıbbi bağlamı veriye dönüştüren "
    "yapay zeka katmanı, (ii) kullanıcının o anki hareket aralığını ölçen "
    "kalibrasyon katmanı ve (iii) bu iki sınırı kesişim alarak birleştiren "
    "hibrit güvenli açı motoru.",
)
add_paragraph(
    doc,
    "Bu doküman, kodu okumadan sistemin nasıl çalıştığını anlamak isteyen "
    "hakemler, değerlendiriciler ve yeni katılan geliştiriciler için hazırlanmıştır. "
    "Her bölüm, karşılığı olan kaynak dosya(lar)a referans verir.",
)

add_callout(
    doc,
    "Tasarım İlkesi",
    "Veri cihazdan çıkmaz. MedGemma yerel GPU’da çalışır, kamera görüntüsü "
    "ve iskelet verisi hiçbir bulut servisine gönderilmez. SQLite veritabanı "
    "yerel diskte durur. Bu tercih, KVKK kapsamındaki özel nitelikli sağlık "
    "verisine donanımsal düzeyde uyum sağlar.",
)

# === 2. Sistem Mimarisi ======================================================
add_heading(doc, "2. Sistem Mimarisi", level=1)
add_paragraph(
    doc,
    "Sistem, tarayıcıda çalışan bir Next.js arayüzü ile yerel makinede çalışan "
    "bir FastAPI arka ucundan oluşur. Kurulum/kalibrasyon/özet gibi işlemler "
    "REST uç noktaları üzerinden yürür; canlı kare analizi ise tek bir WebSocket "
    "üzerinden yapılır. Aşağıdaki akış, bir oturumun yaşam döngüsünü özetler:",
)
add_bullets(
    doc,
    [
        "Tarayıcı egzersiz ve (opsiyonel) sağlık durumu bilgisini POST /api/setup ile gönderir.",
        "Sunucu, şikayet varsa MedGemma’yı çağırır; yoksa egzersizin varsayılan güvenli aralığını kullanır.",
        "3 tekrarlı kalibrasyon verisi POST /api/calibration/{session_id} ile yüklenir; kişisel ROM ile tıbbi sınırın kesişimi saklanır.",
        "WebSocket /ws/session/{session_id} üzerinden her kare JPEG olarak gönderilir; geriye açı, tekrar, uyarı ve form skoru döner.",
        "Oturum sonunda POST /sessions/{id}/summary MedGemma’dan seansa özel kısa bir yorum üretir.",
    ],
)

# === 3. Teknoloji Yığını =====================================================
add_heading(doc, "3. Teknoloji Yığını", level=1)
add_kv_table(
    doc,
    [
        ("Frontend", "Next.js 14, React 18, TypeScript, Tailwind CSS, Zustand (durum), react-webcam"),
        ("Görsel İşleme (Web)", "MediaPipe Pose (tarayıcı tarafı çizim için) + sunucu tarafı BlazePose"),
        ("Backend", "Python 3.10+, FastAPI, Pydantic, SQLModel"),
        ("AI Modeli", "google/medgemma-4b-it (4-bit nicemleme, bitsandbytes)"),
        ("Görsel İşleme (Sunucu)", "MediaPipe Tasks — pose_landmarker_lite.task (VIDEO modu)"),
        ("Veritabanı", "SQLite (yerel, tek dosya)"),
        ("İletişim", "HTTP + WebSocket (saniyede ~25 FPS, sunucu tarafı sınır)"),
    ],
)

# === 4. Proje Dizin Yapısı ===================================================
add_heading(doc, "4. Proje Dizin Yapısı", level=1)
add_code_block(
    doc,
    """FitGuardAI/
├── backend/
│   └── app/
│       ├── main.py                 # FastAPI giriş noktası, HTTP uçları
│       ├── config.py               # Ortam değişkenleri (pydantic-settings)
│       ├── core/                   # Ağır algoritmalar — ML'den bağımsız
│       │   ├── pose_processor.py   # MediaPipe sarmalayıcı
│       │   ├── angle_calculator.py # Vektörel açı hesaplama
│       │   ├── rep_counter.py      # Peak/valley tekrar sayacı
│       │   ├── safety_engine.py    # Güvenli aralık + histerezis + form skoru
│       │   └── asymmetry_detector.py
│       ├── ai/
│       │   ├── medgemma_client.py  # Singleton LLM sarmalayıcı + mock
│       │   ├── prompts.py          # Güvenli aralık & seans özeti promptları
│       │   └── parser.py           # LLM JSON çıktısını güvenle parse eder
│       ├── exercises/              # Egzersiz konfigürasyonları (veri olarak)
│       ├── services/
│       │   ├── session.py          # In-memory oturum deposu
│       │   ├── calibration.py      # 3-tekrar → kişisel ROM → kesişim
│       │   └── history.py          # SQLite kalıcılık
│       ├── models/
│       │   ├── schemas.py          # Pydantic veri sözleşmeleri
│       │   └── db.py               # SQLModel tabloları
│       └── websocket/handlers.py   # Gerçek zamanlı kare işleme
├── frontend/
│   ├── app/                        # Next.js App Router sayfaları
│   ├── components/                 # CameraView, CalibrationWizard, SafetyAlert …
│   └── lib/                        # api.ts, websocket.ts, store.ts, tts.ts
└── scripts/                        # Yardımcı scriptler (sunum üreteci vb.)""",
    caption="Dizin yapısı — özet.",
)

# === 5. Backend ==============================================================
add_heading(doc, "5. Backend Bileşenleri", level=1)

# 5.1 main.py
add_heading(doc, "5.1 Giriş Noktası (main.py)", level=2)
add_paragraph(
    doc,
    "FastAPI uygulaması, lifespan hook’unda SQLite tablolarını oluşturur ve "
    "CORS açarak Next.js geliştirme sunucusuyla konuşur. İş akışının tüm "
    "kilit adımları tek bir dosyada toplanmıştır:",
)
add_bullets(
    doc,
    [
        "GET /health — servis sağlık kontrolü; mock_medgemma bayrağını da döner.",
        "GET /exercises, GET /exercises/{id} — desteklenen egzersizlerin kataloğu.",
        "POST /api/setup — egzersiz + (opsiyonel) şikayeti alır, MedGemma çağrısı sonrası güvenli aralığı döner.",
        "POST /api/calibration/{session_id} — 3 tekrarın min/max örneklerini alır, kişisel ROM ile sentezler.",
        "POST /api/sessions/{id}/end — oturumu kapatır, kalıcılaştırır.",
        "POST /sessions/{id}/summary — MedGemma ile seans özeti üretir ve saklar.",
        "WebSocket /ws/session/{id} — kare bazlı gerçek zamanlı analiz.",
    ],
)
add_code_block(
    doc,
    """@app.post(\"/api/setup\", response_model=SafeRangeResponse)
def setup_session(req: SetupRequest) -> SafeRangeResponse:
    exercise = get_exercise(req.exercise_id)
    condition = (req.condition or \"\").strip()
    if not condition:
        # Şikayet yoksa MedGemma atlanır — varsayılan aralıklar kullanılır.
        payload = SafeRangePayload(modified_ranges=dict(exercise.default_safe_ranges), ...)
    else:
        prompt = build_safe_range_prompt(...)
        llm_output = MedGemmaClient.instance().generate(prompt)
        payload = parse_safe_ranges(llm_output, exercise.default_safe_ranges)
    session = session_store.create(...)
    history_svc.persist_session(session)
    return SafeRangeResponse(...)""",
    caption="main.py — /api/setup uç noktası (özet).",
)

# 5.2 Config
add_heading(doc, "5.2 Konfigürasyon Katmanı", level=2)
add_paragraph(
    doc,
    "config.py, pydantic-settings tabanlıdır ve .env dosyasını okur. Tüm "
    "anahtarlar ortam değişkeniyle üzerine yazılabilir, böylece geliştirme "
    "makinelerinde MOCK_MEDGEMMA=true ile modelsiz çalışmak mümkündür.",
)
add_kv_table(
    doc,
    [
        ("medgemma_model", "google/medgemma-4b-it (varsayılan)"),
        ("mock_medgemma", "true → MedGemma çağrılmaz, deterministik yanıt döner"),
        ("use_quantization", "CUDA varsa 4-bit yükleme (bitsandbytes)"),
        ("db_path", "./fitguard.db — yerel SQLite yolu"),
        ("cors_origins", "http://localhost:3000 (virgülle ayrılabilir liste)"),
    ],
)

# 5.3 Exercises
add_heading(doc, "5.3 Egzersiz Tanımları (exercises/)", level=2)
add_paragraph(
    doc,
    "Her egzersiz bir Pydantic ExerciseConfig nesnesidir. Bu veri-merkezli "
    "tasarım sayesinde yeni bir egzersiz eklemek yalnızca yeni bir .py dosyası "
    "oluşturmayı gerektirir; çekirdek algoritma dokunulmamış kalır. Aşağıda "
    "Squat tanımı özetlenmiştir:",
)
add_code_block(
    doc,
    """SQUAT = ExerciseConfig(
    id=\"squat\", name=\"Squat\", camera_angle=\"side\",
    tracked_joints=[
        JointSpec(name=\"left_knee\", side=\"left\",
                  a=LEFT_HIP, vertex=LEFT_KNEE, c=LEFT_ANKLE),
        JointSpec(name=\"right_knee\", side=\"right\",
                  a=RIGHT_HIP, vertex=RIGHT_KNEE, c=RIGHT_ANKLE),
        JointSpec(name=\"left_hip\", side=\"left\",
                  a=LEFT_SHOULDER, vertex=LEFT_HIP, c=LEFT_KNEE),
    ],
    default_safe_ranges={
        \"left_knee\":  (70.0, 175.0),
        \"right_knee\": (70.0, 175.0),
        \"left_hip\":   (55.0, 175.0),
    },
    rep_detection=RepDetectionConfig(joint_name=\"right_knee\", ...),
    symmetry_check=True,
)""",
    caption="exercises/squat.py — egzersiz konfigürasyonu.",
)
add_paragraph(
    doc,
    "JointSpec, üç MediaPipe landmark indeksini saklar: a (dış), vertex "
    "(açının ölçüldüğü tepe) ve c (dış). symmetry_check=True, sağ/sol "
    "eklemlerin otomatik asimetri kontrolüne alınacağını belirtir. "
    "Varsayılan aralıklar, MedGemma kullanılmadığı durumlarda da tüm "
    "sistemin çalışabilmesi için muhafazakâr seçilmiştir.",
)

# 5.4 Core
add_heading(doc, "5.4 Çekirdek Modüller (core/)", level=2)

add_heading(doc, "pose_processor.py — MediaPipe Sarmalayıcı", level=3)
add_paragraph(
    doc,
    "MediaPipe Tasks API (0.10.x+) ile VIDEO çalışma kipinde çalışır. İlk "
    "kullanımda ~3 MB’lık pose_landmarker_lite.task modelini yerel diske "
    "indirir. Her kare için 33 landmark noktası, görünürlük skoru ile birlikte "
    "döner. Ortalama görünürlük belirlenen eşiği aşmıyorsa kare ‘görünmez’ "
    "olarak işaretlenir ve alt katmanlara geçilmez — bu, sessizce yanlış "
    "analiz yapılmasını engeller.",
)

add_heading(doc, "angle_calculator.py — Vektörel Açı", level=3)
add_paragraph(
    doc,
    "Üç nokta arasındaki açı, iç çarpım (dot product) formülü ile hesaplanır. "
    "2D veya 3D noktalar kabul edilir; herhangi bir vektör sıfır uzunluklu "
    "ise NaN döner ve çağıran taraf bu kareyi atar.",
)
add_code_block(
    doc,
    """def calculate_angle(a, b, c) -> float:
    ba = a - b
    bc = c - b
    if np.linalg.norm(ba) == 0 or np.linalg.norm(bc) == 0:
        return math.nan
    cos_theta = np.dot(ba, bc) / (np.linalg.norm(ba) * np.linalg.norm(bc))
    cos_theta = np.clip(cos_theta, -1.0, 1.0)
    return math.degrees(math.acos(cos_theta))""",
    caption="core/angle_calculator.py — özet.",
)

add_heading(doc, "rep_counter.py — Tekrar Sayacı", level=3)
add_paragraph(
    doc,
    "Sinyalden bağımsız (peak/valley) bir yaklaşım kullanır: kullanıcının "
    "mutlak ROM’u bilinmese bile, yerel tepe ile yerel çukur arasındaki "
    "genlik min_range_deg eşiğini (varsayılan 20°) aşıyorsa tekrar sayılır. "
    "Jitter’a karşı hysteresis_deg (varsayılan 4°) bir tampon uygular; "
    "yön değişimi ancak bu tampon aşıldığında tanınır. Tekrar emiti sonrası "
    "hem _last_peak hem de _last_valley sıfırlanır — yeni tekrar için tam "
    "bir döngü şarttır.",
)
add_code_block(
    doc,
    """if self._direction >= 0:
    if angle >= self._extreme_hi:
        self._extreme_hi = angle
    elif angle < self._extreme_hi - self.hysteresis_deg:
        self._last_peak = self._extreme_hi
        self._direction = -1
        if self._last_valley is not None:
            rng = self._last_peak - self._last_valley
            if rng >= self.min_range_deg:
                produced = self._emit_rep(self._last_valley, self._last_peak)
                self._last_peak = None
                self._last_valley = None
# falling dalı simetriktir.""",
    caption="core/rep_counter.py — peak/valley algoritması.",
)
add_paragraph(
    doc,
    "_score_tempo, tekrar süresini 3000 ms’ye normalize eder: her 30 ms sapma "
    "1 puan düşürür, negatif değer sıfıra kırpılır. Böylece hem çok hızlı hem "
    "çok yavaş tekrarlar puansal olarak cezalandırılır.",
)

add_heading(doc, "safety_engine.py — Güvenlik Motoru", level=3)
add_paragraph(
    doc,
    "Her eklem için (min, max) güvenli aralığa bakar ve iki katmanlı uyarı "
    "üretir: aralığın dışındaysa ‘danger’, aralığa yaklaşıyorsa "
    "(warning_margin_deg içinde, varsayılan 5°) ‘warning’. Hysteresis mantığı, "
    "bir eklemin güvensiz duruma düştükten sonra sınıra geri dönse bile "
    "hysteresis_deg (5°) kadar içeri girene dek ‘danger’ kalmasını sağlar — "
    "sınır etrafındaki titreme spam’ini engeller.",
)
add_code_block(
    doc,
    """def check(self, angles):
    for joint, angle in angles.items():
        lo, hi = self.safe_ranges[joint]
        was_violating = self._violating.get(joint, False)
        if angle < lo or angle > hi:
            is_violating = True
        elif was_violating and (angle < lo + self.hysteresis_deg
                                or angle > hi - self.hysteresis_deg):
            is_violating = True  # histerezis tamponunda
        else:
            is_violating = False
        self._violating[joint] = is_violating
        # is_violating → 'danger'; sınıra yakın → 'warning'.""",
    caption="core/safety_engine.py — kontrol döngüsü.",
)
add_paragraph(
    doc,
    "form_score, kare bazlı bir skor üretir: 100’den başlar ve aralık dışına "
    "çıkan her eklem için min(30, sapma × 2) puan düşürür; sonuç 0’ın altına "
    "inmez. Böylece küçük sapmalar lineer cezalandırılırken her tek eklem "
    "toplam skoru en fazla 30 puan etkiler.",
)

add_heading(doc, "asymmetry_detector.py — Sağ/Sol Farkı", level=3)
add_paragraph(
    doc,
    "tracked_joints içindeki ‘left_*’ / ‘right_*’ çiftlerini tarar ve ikisi de "
    "mevcutsa |sol - sağ| farkını döndürür. WebSocket işleyicisinde bu değer "
    "10°’yi aşarsa ‘warning’ seviyeli bir ‘asymmetry’ uyarısı frame yanıtına "
    "eklenir.",
)

# 5.5 AI
add_heading(doc, "5.5 Yapay Zeka Katmanı (ai/)", level=2)

add_heading(doc, "medgemma_client.py — Singleton LLM Sarmalayıcısı", level=3)
add_paragraph(
    doc,
    "MedGemmaClient, thread-safe bir singleton’dır. Model yalnızca ilk çağrıda "
    "yüklenir (lazy load). CUDA varsa 4-bit nicemleme ile ~5 GB VRAM’e sığar, "
    "yoksa MPS veya CPU’ya düşer. MOCK_MEDGEMMA=true iken hiçbir model "
    "yüklenmez ve deterministik bir örnek yanıt üretilir — bu, geliştirme "
    "ve CI ortamlarında büyük önem taşır.",
)
add_code_block(
    doc,
    """messages = [{\"role\": \"user\", \"content\": prompt}]
inputs = self._tokenizer.apply_chat_template(
    messages, add_generation_prompt=True,
    tokenize=True, return_dict=True, return_tensors=\"pt\",
).to(self._device)
with torch.no_grad():
    output_ids = self._model.generate(**inputs, max_new_tokens=512, do_sample=False)
generated = output_ids[0, inputs[\"input_ids\"].shape[-1]:]
return self._tokenizer.decode(generated, skip_special_tokens=True)""",
    caption="ai/medgemma_client.py — üretim (generate) özet.",
)

add_heading(doc, "prompts.py ve parser.py", level=3)
add_paragraph(
    doc,
    "prompts.py iki şablon tutar: güvenli aralık sorgusu (SAFE_RANGE_PROMPT) "
    "ve seans özeti (SESSION_SUMMARY_PROMPT). İlki, modele ‘yalnızca belirli "
    "şemada JSON döndür’ komutu verir. parser.py ise bu JSON’u savunmacı "
    "şekilde işler: kod bloğu işareti olsa bile regex ile çıkarır, tipleri "
    "zorlar (0–180° aralığı), geçersiz değerleri atlar ve — en kritik ilke — "
    "modelin ürettiği aralıkları yalnızca DARALTMAYA izin verir (genişletemez).",
)
add_code_block(
    doc,
    """for joint, rng in modified.items():
    clean = _sanitize_range(rng)
    if clean is None:
        continue
    # Yalnızca daralt — güvenlik ilkesi.
    d_lo, d_hi = defaults.get(joint, clean)
    payload.modified_ranges[joint] = (max(clean[0], d_lo), min(clean[1], d_hi))""",
    caption="ai/parser.py — biyomekanik doğrulama katmanı.",
)

# 5.6 Services
add_heading(doc, "5.6 Servisler (services/)", level=2)

add_heading(doc, "session.py — In-Memory Oturum Deposu", level=3)
add_paragraph(
    doc,
    "Session dataclass’ı bir oturumun canlı durumunu taşır: MedGemma’dan "
    "gelen safe_ranges, kişisel ROM, etkin aralıklar, tekrar geçmişi, form "
    "skorları, asimetri örnekleri ve uyarı sayısı. SessionStore sözlük "
    "tabanlı basit bir bellek-içi depodur; uygulama yeniden başlatıldığında "
    "history.py SQLite’tan geri yükleme yapar.",
)

add_heading(doc, "calibration.py — Hibrit Eşik", level=3)
add_paragraph(
    doc,
    "Kullanıcının 3 tekrar boyunca ölçtüğü min/max açılar ortalamaya alınarak "
    "kişisel ROM elde edilir. Ardından, her eklem için MedGemma’nın sınırı "
    "ile bu ROM’un kesişimi alınır — yani etkin aralık her zaman her iki "
    "kısıtın en daralmış halidir. MedGemma’da yer almayan eklemler için "
    "güvenli aralık varsayılandan düşer.",
)
add_code_block(
    doc,
    """personal_rom[j] = (sum(mins)/len(mins), sum(maxs)/len(maxs))
effective[j] = (max(p_lo, s_lo), min(p_hi, s_hi))  # kesişim
# MedGemma'da olmayan eklemler için:
for j, rng in session.safe_ranges.items():
    effective.setdefault(j, rng)
session.personal_rom = personal_rom
session.effective_ranges = effective
session.calibrated = True""",
    caption="services/calibration.py — hibrit sınır hesabı.",
)

add_heading(doc, "history.py — SQLite Kalıcılık", level=3)
add_paragraph(
    doc,
    "SQLModel ile üç tablo yönetir: sessions (oturum meta + özet), angle_logs "
    "(eklem-bazlı kare logları) ve rep_logs (tekrar özetleri). Tümü yerel "
    "fitguard.db dosyasında tutulur; hiçbir uzak bağlantı yapılmaz.",
)

# 5.7 WebSocket
add_heading(doc, "5.7 WebSocket İşleyicisi (websocket/handlers.py)", level=2)
add_paragraph(
    doc,
    "FrameProcessor, her bir WebSocket bağlantısı için kendi PoseProcessor, "
    "SafetyEngine ve RepCounter örneklerini tutar. Bu sayede her oturum "
    "bağımsız bir durum makinesi olur ve paralel kullanıcılar birbirinin "
    "analizine karışmaz. İşleyici şunları yapar:",
)
add_bullets(
    doc,
    [
        "MIN_FRAME_INTERVAL_MS ile kare hızını sunucu tarafında ~25 FPS’a sınırlar (istemci daha hızlı göndermeye çalışsa bile).",
        "Base64 kareyi OpenCV ile çözer, MediaPipe’a RGB olarak verir.",
        "Eklem görünürlüğü MIN_LANDMARK_CONFIDENCE (0.3) altındaysa o eklemi atlar.",
        "Açılar → SafetyEngine.check → uyarılar ve form_score üretir.",
        "Egzersizin rep_detection.joint_name’ine karşılık gelen açıyı RepCounter’a verir.",
        "symmetry_check açıksa compute_asymmetry çağırır; 10°’yi aşan fark için asymmetry uyarısı ekler.",
        "Kontrol mesajları: reset (sayacı sıfırla), stop (kapat).",
    ],
)
add_code_block(
    doc,
    """pose_result = self.pose.process_frame(img)
angles = _compute_angles(self.exercise, pose_result)
alerts = self.safety.check(angles)
form_score = self.safety.form_score(angles)
rep_event = None
joint_for_rep = self.exercise.rep_detection.joint_name
if joint_for_rep in angles:
    rep_event = self.rep.update(angles[joint_for_rep])
asymmetry = compute_asymmetry(self.exercise, angles) if self.exercise.symmetry_check else None
if asymmetry is not None and asymmetry > 10.0:
    alerts.append(Alert(level=\"warning\", joint=\"asymmetry\", ...))""",
    caption="websocket/handlers.py — her kare için yapılan iş.",
)

# 5.8 Data models
add_heading(doc, "5.8 Veri Modelleri ve Kalıcılık (models/)", level=2)
add_paragraph(
    doc,
    "schemas.py, istemci ve sunucu arasındaki sözleşmeyi tanımlar. Öne çıkan tipler:",
)
add_bullets(
    doc,
    [
        "AngleSnapshot: her eklem için açı + in_safe_zone + aralık sınırları.",
        "Alert: level {info, warning, danger} + eklem + mesaj + açı.",
        "RepEvent: rep_number, duration_ms, peak, min, tempo_score.",
        "FrameResponse: WebSocket yanıtı — landmarks, angles, rep_count, phase, alerts, form_score, rep_event, asymmetry.",
        "SafeRangeResponse: /api/setup çıktısı — sessionId + aralıklar + MedGemma reasoning/warnings/alternatives.",
        "CalibrationResult: personal_rom + effective_ranges.",
        "SessionSummary: seans-sonu özet (metin + metrikler).",
    ],
)
add_paragraph(
    doc,
    "db.py, üç SQLModel tablosuyla aynı bilgileri diske yazar. Ayrıca "
    "init_db() uygulama ayağa kalkarken bu şemaları kurar.",
)

# === 6. Frontend =============================================================
add_heading(doc, "6. Frontend Bileşenleri", level=1)
add_paragraph(
    doc,
    "Frontend, sayfa yönlendirmesi için Next.js 14 App Router, global durum "
    "için Zustand ve kamera erişimi için react-webcam kullanır. İskelet "
    "çizimi tarayıcıda doğrudan canvas üzerinden yapılır; bu, ağ gecikmesine "
    "bakmaksızın akıcı bir deneyim sağlar.",
)
add_kv_table(
    doc,
    [
        ("lib/store.ts", "Zustand durumu: seçili egzersiz, oturum kimliği, aralıklar, kare sonuçları, resetRepsTick."),
        ("lib/websocket.ts", "SessionWebSocket sınıfı — oto-reconnect, tek kanal, tip-kontrol."),
        ("lib/api.ts", "REST çağrıları (/api/setup, /api/calibration, /sessions/…/summary)."),
        ("lib/tts.ts", "Web Speech API üzerinden Türkçe sesli uyarı."),
        ("components/ExerciseSelector", "Egzersiz seçim kartları."),
        ("components/MedicalConditionInput", "Opsiyonel şikayet girişi (Kişisel Sağlık Paneli)."),
        ("components/CalibrationWizard", "3-tekrarlı yönlendirmeli kalibrasyon akışı."),
        ("components/CameraView", "Kare yakalama + iskelet çizimi + FPS limiti."),
        ("components/SafetyAlert", "Renk kodlu uyarı rozetleri, sesli bildirim."),
        ("components/RepCounter / AngleDisplay", "Canlı metrikler."),
        ("components/SessionReport", "Recharts ile zaman-serisi grafikleri, MedGemma özeti."),
    ],
)
add_paragraph(
    doc,
    "Zustand store’unda resetRepsTick deseni önemlidir: alt bileşenler bu "
    "sayacın değişmesine abone olarak kendi iç durumlarını (zamanlayıcılar, "
    "render buffer’ları) temizler; böylece React’in sığ karşılaştırma "
    "tuzaklarına düşmeden deterministik bir ‘reset’ sinyali yayılır.",
)

# === 7. Uçtan Uca Akış =======================================================
add_heading(doc, "7. Uçtan Uca Kullanıcı Akışı", level=1)
add_bullets(
    doc,
    [
        "Egzersiz Seçimi → ExerciseSelector, GET /exercises ile kataloğu çeker.",
        "Sağlık Paneli → MedicalConditionInput opsiyonel şikayet alır.",
        "Kurulum → POST /api/setup; şikayet varsa MedGemma → güvenli aralık JSON’u; yoksa varsayılan.",
        "Kalibrasyon → CalibrationWizard, 3 ağırlıksız tekrar sırasında tarayıcıda lokal min/max toplar ve POST /api/calibration ile yükler.",
        "Kesişim → calibration.apply_calibration, MedGemma sınırı ile kişisel ROM’un kesişimini üretir; effective_ranges artık egzersiz boyunca referanstır.",
        "Canlı Analiz → CameraView her kareyi WebSocket’e gönderir; FrameProcessor açı + tekrar + uyarı + form skoru + asimetri döner.",
        "Seans Sonu → POST /api/sessions/{id}/end ardından POST /sessions/{id}/summary → MedGemma seansa özel metin üretir.",
        "Geçmiş → /history sayfası SQLite’tan okunan oturum rozetlerini gösterir.",
    ],
)

# === 8. KVKK / Güvenlik ======================================================
add_heading(doc, "8. Güvenlik, Gizlilik ve KVKK", level=1)
add_paragraph(
    doc,
    "Tüm veri akışı şu ilkeler çerçevesinde tasarlanmıştır:",
)
add_bullets(
    doc,
    [
        "Görüntü, iskelet verisi ve sağlık bilgisi cihaz dışına çıkmaz.",
        "MedGemma yerel GPU’da çalışır; 4-bit nicemleme ile RTX 4060 sınıfı kartlara sığar.",
        "SQLite veritabanı tek kullanıcı ve tek makine için yerel bir dosyadır.",
        "CORS, yalnızca açıkça izin verilen kaynaklara açıktır (varsayılan: localhost:3000).",
        "Yapay zeka çıktısı yalnızca güvenli aralığı DARALTABİLİR; genişletemez (parser.py ilkesi).",
        "Üretim sürümünden önce ortam değişkenlerindeki Hugging Face token gibi sırlar repodan arındırılmalıdır.",
    ],
)
add_callout(
    doc,
    "Uyum Çerçevesi",
    "6698 sayılı KVKK, GDPR (Madde 9 — özel nitelikli veri), EU AI Act "
    "(sağlık alanında yüksek-risk AI sistemleri için şeffaflık ve insan "
    "gözetimi ilkeleri) ve tıbbi yazılımlar için IEC 62304 prensipleriyle "
    "uyumlu bir mimari benimsenmiştir. Hiçbir verinin buluta çıkmaması, "
    "uyumu donanımsal düzeyde garanti altına alır.",
)

# === 9. Geliştirme ===========================================================
add_heading(doc, "9. Geliştirme ve Çalıştırma", level=1)
add_heading(doc, "Backend", level=2)
add_code_block(
    doc,
    """# İlk kurulum
python -m venv venv
./venv/Scripts/activate   # Windows
pip install -r backend/requirements.txt

# Modelsiz (hızlı) çalıştırma
set MOCK_MEDGEMMA=true
uvicorn backend.app.main:app --reload --port 8000

# Gerçek MedGemma (CUDA + bitsandbytes gerekir)
set MOCK_MEDGEMMA=false
uvicorn backend.app.main:app --port 8000""",
    caption="Backend — kurulum ve çalıştırma.",
)
add_heading(doc, "Frontend", level=2)
add_code_block(
    doc,
    """cd frontend
npm install
npm run dev   # http://localhost:3000""",
    caption="Frontend — Next.js geliştirme sunucusu.",
)

add_paragraph(
    doc,
    "Tipik geliştirme döngüsünde MOCK_MEDGEMMA=true bayrağı GPU gerektirmez; "
    "MedGemma istemcisi deterministik bir JSON döndürdüğü için tüm akış "
    "(kurulum → kalibrasyon → canlı analiz → özet) uçtan uca test edilebilir.",
)

# --- Alt bilgi ---------------------------------------------------------------
doc.add_paragraph()
footer = doc.add_paragraph()
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run(
    "FitGuard AI · TEKNOFEST 2026 · Yazılım Teknik Dokümantasyonu"
)
r.font.name = "Calibri"
r.font.size = Pt(9)
r.italic = True
r.font.color.rgb = SLATE500

# --- Yaz ---------------------------------------------------------------------
OUT = r"C:\Users\Utku35\Desktop\teknofest\Araştırma Projesi\Yeni klasör\sunum\FitGuardAI_Yazilim_Dokumantasyonu.docx"
doc.save(OUT)
print("WROTE:", OUT)
